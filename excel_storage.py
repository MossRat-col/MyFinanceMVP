
from pathlib import Path
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


# Excel-файл будет находиться в папке проекта
EXCEL_FILE = Path(__file__).with_name("transactions.xlsx")

# Названия столбцов
HEADERS = ["ID", "Дата", "Итог за день (UZS)"]


def create_file_if_missing():
    """Создаёт Excel-файл, если его ещё нет."""

    if EXCEL_FILE.exists():
        return

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "История"

    sheet.append(HEADERS)

    # Оформление заголовков
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    sheet.column_dimensions["A"].width = 10
    sheet.column_dimensions["B"].width = 16
    sheet.column_dimensions["C"].width = 25

    workbook.save(EXCEL_FILE)
    workbook.close()


def load_history():
    """Возвращает всю историю из Excel."""

    create_file_if_missing()

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook.active

    history = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        record_id, record_date, daily_total = row

        # Пропускаем пустые строки
        if record_id is None or record_date is None:
            continue

        # Преобразуем дату в единый формат YYYY-MM-DD
        if isinstance(record_date, (date, datetime)):
            record_date = record_date.isoformat()[:10]
        else:
            record_date = str(record_date)

        history.append({
            "id": int(record_id),
            "date": record_date,
            "daily_total": int(daily_total or 0)
        })

    workbook.close()
    return history


def get_daily_total(date_string):
    """Возвращает итог за указанную дату."""

    history = load_history()

    for record in history:
        if record["date"] == date_string:
            return record["daily_total"]

    # Если такой даты ещё нет, итог равен нулю
    return 0


def save_daily_total(date_string, amount):
    """Добавляет новую дату или обновляет существующую."""

    create_file_if_missing()

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook.active

    # Ищем дату, которая уже есть в таблице
    for row_number in range(2, sheet.max_row + 1):
        cell_value = sheet.cell(row=row_number, column=2).value

        if isinstance(cell_value, (date, datetime)):
            cell_date = cell_value.isoformat()[:10]
        else:
            cell_date = str(cell_value) if cell_value is not None else ""

        if cell_date == date_string:
            # ID сохраняется, изменяется только дневной итог
            sheet.cell(row=row_number, column=3).value = int(amount)

            workbook.save(EXCEL_FILE)
            workbook.close()
            return

    # Если даты нет, создаём новую запись
    existing_ids = [
        sheet.cell(row=row_number, column=1).value
        for row_number in range(2, sheet.max_row + 1)
    ]

    valid_ids = [
        int(record_id)
        for record_id in existing_ids
        if record_id is not None
    ]

    new_id = max(valid_ids, default=0) + 1

    sheet.append([new_id, date_string, int(amount)])

    workbook.save(EXCEL_FILE)
    workbook.close()

#ААААААААААААААААА. Каким образом я это понимаю????
#добавляю новую функцию "Открывания таблицы", то есть оригинальный xslsx файл.


def get_file_path():
    """Возвращает путь к Excel-файлу."""
    create_file_if_missing()
    return EXCEL_FILE.resolve()

#Сброс данных(То есть удаление excel файла и создание нового)

def reset_data():
    """Удаляет старые данные и создаёт новый Excel-файл."""
    create_file_if_missing()
    EXCEL_FILE.unlink()
    create_file_if_missing()